"""GSTR-1 versus GSTR-3B reconciliation logic."""

import calendar
import re
from typing import Any, Dict, List, Optional, Tuple

import pandas as pd

HEADS = ["taxable", "igst", "cgst", "sgst", "cess"]
HEAD_LABELS = {
    "taxable": "Taxable value",
    "igst": "IGST",
    "cgst": "CGST",
    "sgst": "SGST/UTGST",
    "cess": "Cess",
}

G1_COLS = {
    "taxable": "Value (₹)",
    "igst": "Integrated Tax (₹)",
    "cgst": "Central Tax (₹)",
    "sgst": "State/UT Tax (₹)",
    "cess": "Cess (₹)",
}

CDN_SECTIONS = {"9B-CDNR", "9B-CDNUR", "9C-CDNRA", "9C-CDNURA"}

MAPPING: Dict[str, List[Tuple[str, int]]] = {
    "3.1(a)": [
        ("4A", 1), ("5", 1), ("6C", 1), ("7", 1),
        ("9A-B2B", 1), ("9A-B2CL", 1), ("9A-DE", 1), ("10", 1),
        ("9B-CDNR", 1), ("9B-CDNUR", 1), ("9C-CDNRA", 1), ("9C-CDNURA", 1),
        ("11A", 1), ("11B", -1), ("11A-Amend", 1), ("11B-Amend", -1),
    ],
    "3.1(b)": [("6A", 1), ("6B", 1), ("9A-EXP", 1), ("9A-SEZ", 1)],
    "3.1(c)": [("8", 1)],
}
ROW_LABELS = {
    "3.1(a)": "3.1(a) Outward taxable supplies",
    "3.1(b)": "3.1(b) Zero rated supplies",
    "3.1(c)": "3.1(c) Nil rated / exempted",
}
ROW_PATTERNS = {
    "3.1(a)": r"outward taxable supplies.*other than zero",
    "3.1(b)": r"outward taxable supplies.*zero",
    "3.1(c)": r"other outward supplies.*nil",
    "3.1(d)": r"inward supplies.*reverse charge",
    "3.1(e)": r"non.?gst outward",
}

_MONTHS = {month.lower(): index for index, month in enumerate(calendar.month_name) if month}
_MONTHS.update({month.lower(): index for index, month in enumerate(calendar.month_abbr) if month})


def _num(value: Any) -> float:
    if value is None:
        return 0.0
    text = str(value).replace(",", "").replace("₹", "").strip()
    if text.lower() in ("", "-", "--", "nan", "none"):
        return 0.0
    negative = text.startswith("(") and text.endswith(")")
    try:
        number = float(text.strip("()"))
    except ValueError:
        return 0.0
    return -number if negative else number


def period_key(metadata: Dict[str, str]) -> Optional[Tuple[int, int]]:
    """Return (financial-year start, calendar month) for supported metadata."""
    month = _MONTHS.get(str(metadata.get("period", "")).strip().lower())
    year_match = re.search(r"\d{4}", str(metadata.get("year", "")))
    if not month or not year_match:
        return None

    year = int(year_match.group())
    is_fy = bool(re.search(r"\d{4}\s*-\s*\d{2,4}", str(metadata.get("year", ""))))
    financial_year = year if is_fy or month >= 4 else year - 1
    return financial_year, month


def _period_label(key: Tuple[int, int]) -> str:
    financial_year, month = key
    calendar_year = financial_year if month >= 4 else financial_year + 1
    return f"{calendar.month_abbr[month]} {calendar_year}"


def _is_total_row(row: Dict[str, Any]) -> bool:
    return re.match(r"\s*total\b", str(row.get("Description", "")), re.IGNORECASE) is not None


def _cdn_sign(row: Dict[str, Any], values: Dict[str, float]) -> Dict[str, float]:
    text = f"{row.get('Document Type', '')} {row.get('Description', '')}".lower()
    if "net" in text:
        return values
    if "credit" in text:
        return {key: -abs(value) for key, value in values.items()}
    if "debit" in text:
        return {key: abs(value) for key, value in values.items()}
    return values


def g1_section_totals(ext: Dict[str, Any], section_id: str) -> Dict[str, float]:
    totals = dict.fromkeys(HEADS, 0.0)
    section = ext.get("sections", {}).get(section_id)
    if not section:
        return totals

    rows = section.get("rows", [])
    printed_totals = [row for row in rows if _is_total_row(row)]
    if section_id in CDN_SECTIONS:
        rows_to_use = [row for row in rows if not _is_total_row(row)] or printed_totals
    else:
        rows_to_use = printed_totals[:1] or rows

    for row in rows_to_use:
        values = {head: _num(row.get(G1_COLS[head])) for head in HEADS}
        if section_id in CDN_SECTIONS:
            values = _cdn_sign(row, values)
        for head in HEADS:
            totals[head] += values[head]
    return totals


def g1_mapped(ext: Dict[str, Any]) -> Dict[str, Dict[str, float]]:
    result = {}
    for label, parts in MAPPING.items():
        totals = dict.fromkeys(HEADS, 0.0)
        for section_id, sign in parts:
            section_totals = g1_section_totals(ext, section_id)
            for head in HEADS:
                totals[head] += sign * section_totals[head]
        result[label] = totals
    return result


def _find_3b_columns(columns: List[str]) -> Dict[str, str]:
    result: Dict[str, str] = {}
    for column in columns:
        lowered = column.lower()
        if "taxable" in lowered:
            result.setdefault("taxable", column)
        elif "integrated" in lowered:
            result.setdefault("igst", column)
        elif "central" in lowered:
            result.setdefault("cgst", column)
        elif "state" in lowered or "utgst" in lowered:
            result.setdefault("sgst", column)
        elif "cess" in lowered:
            result.setdefault("cess", column)
    return result


def g3b_table_3_1(ext: Dict[str, Any]) -> Dict[str, Dict[str, float]]:
    result = {key: dict.fromkeys(HEADS, 0.0) for key in ROW_PATTERNS}
    seen = set()
    section = ext.get("sections", {}).get("3.1")
    if not section:
        return result

    for row in section.get("rows", []):
        text = re.sub(r"\s+", " ", " ".join(str(value) for value in row.values()).lower())
        label = next((key for key, pattern in ROW_PATTERNS.items() if re.search(pattern, text)), None)
        if not label or label in seen:
            continue
        seen.add(label)
        columns = _find_3b_columns(list(row.keys()))
        for head in HEADS:
            if head in columns:
                result[label][head] = _num(row[columns[head]])
    return result


def _status(difference: float, tolerance: float) -> str:
    if abs(difference) <= tolerance:
        return "Match"
    return "3B lower than GSTR-1" if difference > 0 else "3B higher than GSTR-1"


def reconcile_period(
    g1_ext: Dict[str, Any],
    g3b_ext: Dict[str, Any],
    tol: float = 1.0,
) -> Tuple[List[Dict[str, Any]], List[str]]:
    key = period_key(g1_ext.get("metadata", {})) or period_key(g3b_ext.get("metadata", {}))
    period = _period_label(key) if key else "Unknown"
    g1 = g1_mapped(g1_ext)
    g3b = g3b_table_3_1(g3b_ext)

    rows = []
    issues = []
    for label in MAPPING:
        for head in HEADS:
            difference = g1[label][head] - g3b[label][head]
            rows.append({
                "Period": period,
                "GSTR-3B row": ROW_LABELS[label],
                "Head": HEAD_LABELS[head],
                "GSTR-1": round(g1[label][head], 2),
                "GSTR-3B": round(g3b[label][head], 2),
                "Difference": round(difference, 2),
                "Status": _status(difference, tol),
            })

    printed = g1_section_totals(g1_ext, "Total")
    if any(printed[head] for head in HEADS):
        for head in ("igst", "cgst", "sgst", "cess"):
            mapped_tax = sum(g1[label][head] for label in ("3.1(a)", "3.1(b)"))
            if abs(mapped_tax - printed[head]) > tol:
                issues.append(
                    f"{period}: GSTR-1 mapped {HEAD_LABELS[head]} {mapped_tax:,.2f} != printed "
                    f"Total Liability {printed[head]:,.2f}. Check extraction or MAPPING."
                )

    g1_gstin = g1_ext.get("metadata", {}).get("gstin")
    g3b_gstin = g3b_ext.get("metadata", {}).get("gstin")
    if g1_gstin and g3b_gstin and g1_gstin != g3b_gstin:
        issues.append(f"{period}: GSTIN mismatch (GSTR-1 {g1_gstin} vs GSTR-3B {g3b_gstin}).")
    return rows, issues


def reconcile(
    g1_extractions: List[Dict[str, Any]],
    g3b_extractions: List[Dict[str, Any]],
    tol: float = 1.0,
) -> Tuple[pd.DataFrame, List[str]]:
    """Pair GSTR-1 and GSTR-3B extractions by period and reconcile each pair."""
    g1_by, g3b_by, issues = {}, {}, []
    for extractions, store, name in (
        (g1_extractions, g1_by, "GSTR-1"),
        (g3b_extractions, g3b_by, "GSTR-3B"),
    ):
        for extraction in extractions:
            key = period_key(extraction.get("metadata", {}))
            if key is None:
                issues.append(f"{name}: could not read period from a PDF (metadata: {extraction.get('metadata')}).")
            elif key in store:
                issues.append(f"{name}: duplicate filing period {_period_label(key)}; using the last file supplied.")
            store[key] = extraction

    all_rows = []
    for key in sorted(set(g1_by) | set(g3b_by)):
        if key not in g1_by:
            issues.append(f"{_period_label(key)}: GSTR-3B uploaded but no GSTR-1.")
        elif key not in g3b_by:
            issues.append(f"{_period_label(key)}: GSTR-1 uploaded but no GSTR-3B.")
        else:
            rows, period_issues = reconcile_period(g1_by[key], g3b_by[key], tol)
            all_rows.extend(rows)
            issues.extend(period_issues)
    return pd.DataFrame(all_rows), issues


def build_reco_excel(
    df: pd.DataFrame,
    issues: List[str],
    output_path: str,
    tol: float = 1.0,
) -> str:
    from openpyxl import Workbook
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Reconciliation"
    worksheet["A1"], worksheet["B1"] = "Tolerance (₹)", tol
    worksheet["A1"].font = Font(bold=True)

    headers = ["Period", "GSTR-3B row", "Head", "GSTR-1", "GSTR-3B", "Difference", "Status"]
    header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
    for column_index, header in enumerate(headers, 1):
        cell = worksheet.cell(row=3, column=column_index, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for row_index, row in enumerate(df.itertuples(index=False), start=4):
        worksheet.cell(row=row_index, column=1, value=row[0])
        worksheet.cell(row=row_index, column=2, value=row[1])
        worksheet.cell(row=row_index, column=3, value=row[2])
        worksheet.cell(row=row_index, column=4, value=row[3])
        worksheet.cell(row=row_index, column=5, value=row[4])
        worksheet.cell(row=row_index, column=6, value=f"=D{row_index}-E{row_index}")
        worksheet.cell(
            row=row_index,
            column=7,
            value=(
                f'=IF(ABS(F{row_index})<=$B$1,"Match",'
                f'IF(F{row_index}>0,"3B lower than GSTR-1","3B higher than GSTR-1"))'
            ),
        )
        for column_index in (4, 5, 6):
            worksheet.cell(row=row_index, column=column_index).number_format = "#,##0.00"

    last_row = 3 + len(df)
    if last_row >= 4:
        worksheet.conditional_formatting.add(
            f"A4:G{last_row}",
            FormulaRule(
                formula=['$G4<>"Match"'],
                fill=PatternFill(start_color="FDE2E2", end_color="FDE2E2", fill_type="solid"),
            ),
        )
    for column_index, width in enumerate([12, 34, 16, 16, 16, 16, 24], 1):
        worksheet.column_dimensions[get_column_letter(column_index)].width = width
    worksheet.freeze_panes = "A4"

    issues_worksheet = workbook.create_sheet("Issues")
    issues_worksheet["A1"] = "Extraction / pairing issues"
    issues_worksheet["A1"].font = Font(bold=True)
    for row_index, message in enumerate(issues or ["None"], start=2):
        issues_worksheet.cell(row=row_index, column=1, value=message)
    issues_worksheet.column_dimensions["A"].width = 120

    workbook.save(output_path)
    return output_path
